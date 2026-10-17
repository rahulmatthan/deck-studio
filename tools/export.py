#!/usr/bin/env python3
"""export — turn a deck folder into finished files, and publish share links.

  export.py <deck_dir> [--formats html,pdf,pptx] [--dest D]   write the finished files
  export.py publish <deck_dir> <today|24|72|168|x> [--dry-run] a present-only link with an expiry
  export.py published | revoke <name>                          list / remove published links

html  one self-contained file: images, fonts and chart/QR libraries inlined, so it presents offline
pdf   headless Chrome print, one fully built slide per page
pptx  for Google Slides: each slide is a picture of the design with the text taken off it, and the
      text laid back on top as real, editable text boxes in the same Google fonts

Everything here is stdlib except the pptx writer, which re-runs this file under
`uv run --with python-pptx --with pillow` (nothing is installed globally).
"""
import base64, datetime, hashlib, html, json, mimetypes, os, re, shutil, subprocess, sys, tempfile
import urllib.parse, urllib.request

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import deck

CACHE = os.path.expanduser('~/.cache/deck-studio')
UA = ('Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 '
      '(KHTML, like Gecko) Chrome/129.0.0.0 Safari/537.36')
LIBS = {'chart': 'https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js',
        'qr': 'https://cdn.jsdelivr.net/npm/qrcodejs@1.0.0/qrcode.min.js'}
GH = {'owner': 'rahulmatthan', 'repo': 'deck-studio', 'branch': 'main', 'site': 'https://exmachina.in/deck-studio'}
# expiry gate for published decks — verbatim from the old studio (index.html #tpl-gate)
GATE = ("(function(){try{var m=document.querySelector('meta[name=deck-expires]');if(!m)return;var v=m.content;"
        "if(v&&v!=='never'&&Date.now()>parseInt(v,10)){document.documentElement.innerHTML='<head><meta charset=utf-8>"
        "<title>Link expired</title></head><body style=\"font-family:system-ui,-apple-system,sans-serif;display:flex;"
        "align-items:center;justify-content:center;height:100vh;margin:0;background:#0B1623;color:#F4F6F8\">"
        "<div style=\"text-align:center\"><div style=\"font-size:22px;font-weight:600\">This link has expired.</div>"
        "<div style=\"opacity:.6;margin-top:8px;font-size:14px\">Please ask the sender for a fresh link.</div></div>"
        "</body>';}}catch(e){}})();")
CLOSE_SCRIPT = '</scr' + 'ipt>'


# ---------- small helpers ----------

def fetch(url, headers=None, binary=True):
    """GET with a disk cache, so re-exports work offline once a deck has been exported."""
    os.makedirs(CACHE, exist_ok=True)
    key = hashlib.sha1((url + json.dumps(headers or {})).encode()).hexdigest()
    p = os.path.join(CACHE, key)
    if not os.path.exists(p):
        req = urllib.request.Request(url, headers=headers or {'User-Agent': UA})
        with urllib.request.urlopen(req, timeout=30) as r:
            data = r.read()
        open(p + '.part', 'wb').write(data)
        os.replace(p + '.part', p)
    data = open(p, 'rb').read()
    return data if binary else data.decode('utf-8')


def data_uri(path_or_bytes, mime=None):
    if isinstance(path_or_bytes, bytes):
        b = path_or_bytes
    else:
        b = open(path_or_bytes, 'rb').read()
        mime = mime or mimetypes.guess_type(path_or_bytes)[0]
    return f'data:{mime or "application/octet-stream"};base64,{base64.b64encode(b).decode()}'


def title_of(src):
    m = re.search(r'<title>(.*?)</title>', src, re.S)
    return html.unescape(m.group(1)).strip() if m else 'Presentation'


def safe_name(s):
    s = re.sub(r'[\/\\:*?"<>|\n\r\t]', '-', s)
    return re.sub(r'\s+', ' ', s).strip() or 'Presentation'


def base_name(src):
    return f'{datetime.date.today().isoformat()} — {safe_name(title_of(src))}'


def n_slides(src):
    """slides in the stage only (the engine's own comments mention <section> too)"""
    body = src.split('<div id="stage">', 1)[-1].split('<script id="deck-js">', 1)[0]
    return len(re.findall(r'<section\s+class="slide[ "]', re.sub(r'<!--.*?-->', '', body, flags=re.S)))


def inject(src, head='', body=''):
    if head:
        src = src.replace('</head>', head + '\n</head>', 1)
    if body:
        i = src.rfind('</body>')
        src = src[:i] + body + '\n' + src[i:]
    return src


# ---------- self-contained HTML ----------

def font_css(url):
    """Google Fonts css2 → @font-face blocks with woff2 inlined (latin + latin-ext only).
    Variable fonts list one file under several weights; those merge into one weight-range block."""
    css = fetch(url, binary=False)
    groups = {}
    for label, block in re.findall(r'/\*\s*([\w-]+)\s*\*/\s*(@font-face\s*{[^}]*})', css):
        if label not in ('latin', 'latin-ext'):
            continue
        f = lambda k: (re.search(k + r'\s*:\s*([^;]+);', block) or [None, ''])[1].strip()
        src = re.search(r'url\(([^)]+)\)', block).group(1)
        key = (f('font-family'), f('font-style'), src, f('unicode-range'))
        groups.setdefault(key, []).append(int(f('font-weight') or 400))
    out = []
    for (fam, style, src, rng), ws in groups.items():
        w = str(ws[0]) if len(set(ws)) == 1 else f'{min(ws)} {max(ws)}'
        out.append(f"@font-face{{font-family:{fam};font-style:{style};font-weight:{w};font-display:block;"
                   f"src:url({data_uri(fetch(src), 'font/woff2')}) format('woff2');unicode-range:{rng}}}")
    return '\n'.join(out)


def inline_fonts(src):
    m = re.search(r'<!--deck:fonts-->(.*?)<!--/deck:fonts-->', src, re.S)
    if not m:
        return src
    urls = re.findall(r'href="(https://fonts\.googleapis\.com/css2[^"]+)"', m.group(1))
    try:
        css = '\n'.join(font_css(html.unescape(u)) for u in urls)
    except OSError as e:
        print(f'warning: could not fetch fonts ({e}); the file will need the internet for its fonts', file=sys.stderr)
        return src
    block = f'<!--deck:fonts-->\n<style id="deck-fonts">\n{css}\n</style>\n<!--/deck:fonts-->'
    return src[:m.start()] + block + src[m.end():]


def inline_images(src, d):
    def local(ref):
        ref = html.unescape(ref.strip().strip('\'"'))
        if re.match(r'^(data:|https?:|//|#)', ref):
            return None
        p = os.path.join(d, urllib.parse.unquote(ref.split('?')[0].split('#')[0]))
        return p if os.path.isfile(p) else None

    def attr(m):
        p = local(m.group(2))
        return f'{m.group(1)}"{data_uri(p)}"' if p else m.group(0)

    def url(m):
        p = local(m.group(1))
        return f'url("{data_uri(p)}")' if p else m.group(0)
    src = re.sub(r'(\s(?:src|poster)=)"([^"]+)"', attr, src)
    return re.sub(r'url\(\s*([^)]+?)\s*\)', url, src)


def inline_libs(src):
    """Chart.js / qrcodejs inline when the deck uses them; the engine skips its CDN load if they exist."""
    need = []
    if re.search(r'<canvas[^>]*data-chart', src):
        need.append('chart')
    if re.search(r'id="qr"', src):
        need.append('qr')
    tags = ''
    for k in need:
        try:
            js = fetch(LIBS[k], binary=False).replace('</script', '<\\/script')
            tags += f'<script id="deck-lib-{k}">\n{js}\n{CLOSE_SCRIPT}\n'
        except OSError as e:
            print(f'warning: could not fetch {k} library ({e}); it will load from the CDN', file=sys.stderr)
    return src.replace('<script id="deck-js">', tags + '<script id="deck-js">', 1) if tags else src


def single_html(deck_dir, images=True):
    path = deck.deck_path(deck_dir)
    src = open(path, encoding='utf-8').read()
    src = inline_fonts(src)
    if images:
        src = inline_images(src, os.path.dirname(os.path.abspath(path)))
    return inline_libs(src)


# ---------- PDF ----------

# visit every slide once so charts and QR codes draw, then go back to the first
PRIME = ('<script>addEventListener("load",function(){if(window.Chart)Chart.defaults.animation=false;var n=document.querySelectorAll("#stage section.slide").length;'
         'for(var k=0;k<n;k++)postMessage({deck:"goto",i:k},"*");postMessage({deck:"goto",i:0},"*");});' + CLOSE_SCRIPT)
STILL = '<style>*,*::before,*::after{animation:none!important;transition:none!important}</style>'


def still_charts(src):
    """charts drawn at their final state, before the engine draws the first slide"""
    return src.replace('<script id="deck-js">', '<script>if(window.Chart)Chart.defaults.animation=false;'
                       + CLOSE_SCRIPT + '\n<script id="deck-js">', 1)


def chrome(*args, budget=8000):
    return subprocess.run([deck.chrome_bin(), '--headless=new', '--disable-gpu', '--hide-scrollbars',
                           '--no-first-run', '--no-default-browser-check', f'--virtual-time-budget={budget}', *args],
                          capture_output=True, text=True, timeout=180)


def pdf(deck_dir, out):
    tmp = tempfile.mkdtemp(prefix='deck-pdf-')
    try:
        page = os.path.join(tmp, 'deck.html')
        open(page, 'w', encoding='utf-8').write(inject(still_charts(single_html(deck_dir)), head=STILL, body=PRIME))
        chrome('--no-pdf-header-footer', f'--print-to-pdf={out}', 'file://' + page, budget=10000)
        if not os.path.exists(out):
            raise RuntimeError('Chrome did not write the PDF')
        pages = len(re.findall(rb'/Type\s*/Page[^s]', open(out, 'rb').read()))
        n = n_slides(open(deck.deck_path(deck_dir), encoding='utf-8').read())
        if pages != n:
            print(f'warning: PDF has {pages} pages for {n} slides', file=sys.stderr)
        return out
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------- PPTX: measure in Chrome, write with python-pptx ----------

MEASURE_JS = r'''
(function(){
  function px(v){return parseFloat(v)||0;}
  function fam(f){return (f||'').split(',')[0].replace(/["']/g,'').trim();}
  var SKIP='.ribbon,.deck-logo,aside.notes,script,style,canvas,#qr,svg';
  function direct(el){for(var c=el.firstChild;c;c=c.nextSibling)if(c.nodeType===3&&/\S/.test(c.nodeValue))return true;return false;}
  function blocks(el,out){for(var c=el.firstElementChild;c;c=c.nextElementSibling){
    if(c.matches(SKIP))continue;var cs=getComputedStyle(c);if(cs.display==='none'||cs.visibility==='hidden')continue;
    if(direct(c))out.push(c);else blocks(c,out);}return out;}
  function struck(n,stop){for(var e=n;e&&e!==stop.parentElement;e=e.parentElement){if(/line-through/.test(getComputedStyle(e).textDecorationLine))return true;}return false;}
  function runs(b){var out=[];
    (function walk(n){for(var c=n.firstChild;c;c=c.nextSibling){
      if(c.nodeType===3){var e=c.parentElement,cs=getComputedStyle(e),t=c.nodeValue.replace(/\s+/g,' ');if(!t)continue;
        if(cs.textTransform==='uppercase')t=t.toUpperCase();else if(cs.textTransform==='lowercase')t=t.toLowerCase();
        out.push({t:t,color:cs.color,w:+cs.fontWeight||400,i:cs.fontStyle==='italic',s:struck(e,b),font:fam(cs.fontFamily),size:px(cs.fontSize),
          ls:cs.letterSpacing==='normal'?0:px(cs.letterSpacing)});}
      else if(c.nodeType===1){if(c.tagName==='BR')out.push({br:true});else if(!c.matches(SKIP)&&getComputedStyle(c).display!=='none')walk(c);}}})(b);
    /* collapse whitespace across runs, trim the ends */
    var prevSpace=true;out.forEach(function(r){if(r.br){prevSpace=true;return;}if(prevSpace)r.t=r.t.replace(/^ /,'');prevSpace=/ $/.test(r.t);});
    for(var k=out.length-1;k>=0;k--){if(out[k].br)continue;out[k].t=out[k].t.replace(/ $/,'');if(out[k].t)break;}
    return out.filter(function(r){return r.br||r.t;});}
  function measure(){
    var slides=[].slice.call(document.querySelectorAll('#stage section.slide')),res=[];
    slides.forEach(function(s){
      slides.forEach(function(o){o.classList.toggle('active',o===s);});
      s.querySelectorAll('.frag').forEach(function(f){f.classList.add('shown');});
      var R=s.getBoundingClientRect(),k=1280/R.width;
      function rect(e){var r=e.getBoundingClientRect();return {x:(r.left-R.left)*k,y:(r.top-R.top)*k,w:r.width*k,h:r.height*k};}
      var bg=getComputedStyle(s).backgroundColor;
      var tb=blocks(s,[]).map(function(b){var cs=getComputedStyle(b),fs=px(cs.fontSize);
        return {rect:rect(b),font:fam(cs.fontFamily),size:fs,w:+cs.fontWeight||400,align:cs.textAlign,
          lh:cs.lineHeight==='normal'?fs*1.2:px(cs.lineHeight),runs:runs(b)};}).filter(function(b){return b.runs.length;});
      var imgs=s.classList.contains('bleed')?[]:[].slice.call(s.querySelectorAll('img')).filter(function(m){return !m.closest('.deck-logo');}).map(function(m){
        var cs=getComputedStyle(m);return {rect:rect(m),src:m.currentSrc||m.src,nw:m.naturalWidth,nh:m.naturalHeight,fit:cs.objectFit,pos:cs.objectPosition};});
      var nt=s.querySelector('aside.notes');
      res.push({id:s.dataset.id,bg:bg,blocks:tb,images:imgs,notes:nt?nt.innerText||nt.textContent:''});
    });
    var o=document.createElement('script');o.type='application/json';o.id='measure-out';
    o.textContent=JSON.stringify(res).replace(/</g,'\\u003c');document.body.appendChild(o);
  }
  addEventListener('load',function(){(document.fonts?document.fonts.ready:Promise.resolve()).then(function(){setTimeout(measure,300);});});
})();
'''
# a slide with only its glyphs taken off: keeps backgrounds, highlight blocks, bars, ribbon, logo, charts
BG_CSS = ('<style>#stage section.slide *{-webkit-text-fill-color:transparent!important;'
          'text-decoration-color:transparent!important;text-shadow:none!important;caret-color:transparent!important}'
          '#stage section.slide:not(.bleed) img:not(.deck-logo img){visibility:hidden!important}</style>')


def measure(page):
    r = chrome('--window-size=1280,720', '--dump-dom', 'file://' + page, budget=6000)
    m = re.search(r'<script type="application/json" id="measure-out">(.*?)</script>', r.stdout, re.S)
    if not m:
        raise RuntimeError('could not measure the slides (Chrome output had no measurements)')
    return json.loads(m.group(1))


def pptx(deck_dir, out):
    path = os.path.abspath(deck.deck_path(deck_dir))
    d = os.path.dirname(path)
    src = still_charts(single_html(deck_dir, images=False))    # fonts inlined; images stay relative so measured srcs map to files
    mpage, bpage = os.path.join(d, '.export-measure.html'), os.path.join(d, '.export-bg.html')
    work = tempfile.mkdtemp(prefix='deck-pptx-')
    try:
        open(mpage, 'w', encoding='utf-8').write(inject(src, head=STILL, body=f'<script>{MEASURE_JS}{CLOSE_SCRIPT}'))
        open(bpage, 'w', encoding='utf-8').write(inject(src, head=STILL + BG_CSS))
        data = measure(mpage)
        for k in range(1, len(data) + 1):
            chrome('--window-size=1280,720', '--force-device-scale-factor=2',
                   f'--screenshot={os.path.join(work, f"bg{k:02d}.png")}', f'file://{bpage}#{k},bare', budget=4000)
        json.dump(data, open(os.path.join(work, 'measure.json'), 'w'))
        r = subprocess.run(['uv', 'run', '--quiet', '--no-project', '--with', 'python-pptx', '--with', 'pillow',
                            'python3', os.path.abspath(__file__), '_pptx', work, out], capture_output=True, text=True)
        if r.returncode or not os.path.exists(out):
            raise RuntimeError('pptx writer failed:\n' + (r.stderr or r.stdout)[-1500:])
        return out
    finally:
        for p in (mpage, bpage):
            if os.path.exists(p):
                os.remove(p)
        if os.environ.get('DECK_KEEP'):
            print('kept', work, file=sys.stderr)
        else:
            shutil.rmtree(work, ignore_errors=True)


def write_pptx(work, out):
    """Runs under uv with python-pptx + Pillow."""
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.ns import qn
    from PIL import Image

    E = 9525                                    # EMU per canvas px (1280px = 13.333in)
    data = json.load(open(os.path.join(work, 'measure.json')))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * E), Emu(720 * E)
    blank = prs.slide_layouts[6]
    ALIGN = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.JUSTIFY}

    def rgb(c):
        n = [float(x) for x in re.findall(r'[\d.]+', c or '')]
        return RGBColor(*[int(round(v)) for v in n[:3]]) if len(n) >= 3 and (len(n) < 4 or n[3] > 0) else None

    def flat(png):
        """flat slides stay PNG; photographic ones become JPEG"""
        if os.path.getsize(png) <= 600_000:
            return png
        jpg = png[:-4] + '.jpg'
        Image.open(png).convert('RGB').save(jpg, quality=85, optimize=True)
        return jpg

    def local_image(srcurl, k):
        if srcurl.startswith('file://'):
            return urllib.parse.unquote(urllib.parse.urlparse(srcurl).path)
        p = os.path.join(work, f'img{k}')
        if srcurl.startswith('data:'):
            open(p, 'wb').write(base64.b64decode(srcurl.split(',', 1)[1]))
        else:
            open(p, 'wb').write(fetch(srcurl))
        return p

    def place_image(sl, im, k):
        r, p = im['rect'], local_image(im['src'], k)
        if not os.path.exists(p) or not im['nw']:
            return
        nw, nh = im['nw'], im['nh']
        if im['fit'] == 'contain':
            s = min(r['w'] / nw, r['h'] / nh)
            w, h = nw * s, nh * s
            sl.shapes.add_picture(p, Emu(int((r['x'] + (r['w'] - w) / 2) * E)), Emu(int((r['y'] + (r['h'] - h) / 2) * E)),
                                  Emu(int(w * E)), Emu(int(h * E)))
            return
        pic = sl.shapes.add_picture(p, Emu(int(r['x'] * E)), Emu(int(r['y'] * E)), Emu(int(r['w'] * E)), Emu(int(r['h'] * E)))
        if im['fit'] == 'cover':
            s = max(r['w'] / nw, r['h'] / nh)
            dw, dh = nw * s, nh * s
            pos = [float(v[:-1]) / 100 if v.endswith('%') else None for v in (im['pos'] or '50% 50%').split()[:2]]
            px_, py_ = [(v if v is not None else .5) for v in (pos + [.5, .5])[:2]]
            ox, oy = (dw - r['w']) * px_, (dh - r['h']) * py_
            pic.crop_left, pic.crop_right = ox / dw, (dw - r['w'] - ox) / dw
            pic.crop_top, pic.crop_bottom = oy / dh, (dh - r['h'] - oy) / dh

    def text_box(sl, b):
        r, align = b['rect'], b['align']
        single = r['h'] < b['lh'] * 1.5
        grow = r['w'] * .04
        x = r['x'] - (grow / 2 if align == 'center' else grow if align in ('right', 'end') else 0)
        tb = sl.shapes.add_textbox(Emu(int(x * E)), Emu(int(r['y'] * E)), Emu(int((r['w'] + grow) * E)), Emu(int(r['h'] * E)))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = not single
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.auto_size = None
        paras = [[]]
        for run in b['runs']:
            if run.get('br'):
                paras.append([])
            else:
                paras[-1].append(run)
        for n, runs in enumerate(paras):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
            p.line_spacing = Pt(b['lh'] * .75)
            for run in runs:
                rr = p.add_run()
                rr.text = run['t']
                f = rr.font
                f.name, f.size = run['font'] or b['font'], Pt(round(run['size'] * .75 * 2) / 2)
                f.bold, f.italic = run['w'] >= 600, run['i']
                c = rgb(run['color'])
                if c:
                    f.color.rgb = c
                rpr = rr._r.get_or_add_rPr()
                if run['s']:
                    rpr.set('strike', 'sngStrike')
                if run['ls']:
                    rpr.set('spc', str(int(round(run['ls'] * .75 * 100))))
                for tag in ('a:latin', 'a:ea', 'a:cs'):     # same face for every script
                    el = rpr.find(qn(tag))
                    if el is None:
                        el = rpr.makeelement(qn(tag), {})
                        rpr.append(el)
                    el.set('typeface', f.name)

    for k, s in enumerate(data, 1):
        sl = prs.slides.add_slide(blank)
        c = rgb(s['bg'])
        if c:
            sl.background.fill.solid()
            sl.background.fill.fore_color.rgb = c
        bg = os.path.join(work, f'bg{k:02d}.png')
        if os.path.exists(bg):
            sl.shapes.add_picture(flat(bg), 0, 0, prs.slide_width, prs.slide_height)
        for n, im in enumerate(s['images']):
            place_image(sl, im, f'{k}-{n}')
        for b in s['blocks']:
            text_box(sl, b)
        if s['notes'].strip():
            sl.notes_slide.notes_text_frame.text = re.sub(r'[ \t]+', ' ', s['notes']).strip()
    prs.save(out)


# ---------- the export command ----------

def export_dir():
    return os.path.expanduser(os.environ.get('DECK_EXPORT_DIR') or '~/Documents/speaking')


def export(deck_dir, formats=('html', 'pdf', 'pptx'), dest=None):
    dest = os.path.expanduser(dest or export_dir())
    try:
        os.makedirs(dest, exist_ok=True)
        probe = os.path.join(dest, '.deck-write-test')
        open(probe, 'w').close()
        os.remove(probe)
    except PermissionError as e:
        raise PermissionError(f'cannot write to {dest} ({e.strerror}). On the Mac, give the terminal app '
                              f'Files & Folders access to Documents, or set DECK_EXPORT_DIR to another folder.') from e
    src = open(deck.deck_path(deck_dir), encoding='utf-8').read()
    base = os.path.join(dest, base_name(src))
    out = []
    for f in formats:
        if f == 'html':
            open(base + '.html', 'w', encoding='utf-8').write(single_html(deck_dir))
            out.append(base + '.html')
        elif f == 'pdf':
            out.append(pdf(deck_dir, base + '.pdf'))
        elif f == 'pptx':
            out.append(pptx(deck_dir, base + '.pptx'))
        else:
            raise ValueError(f'unknown format "{f}" (html, pdf, pptx)')
    return out


# ---------- publish links (GitHub Pages, public repo deck-studio, folder d/) ----------

def gh(*args, input_file=None):
    cmd = ['gh', 'api', *args] + (['--input', input_file] if input_file else [])
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode:
        raise RuntimeError(f'gh api failed: {(r.stderr or r.stdout).strip()[:300]}')
    return json.loads(r.stdout) if r.stdout.strip() else None


def expiry_epoch(sel):
    if sel == 'x':
        return 'x'
    now = datetime.datetime.now()
    if sel == 'today':
        return str(int(now.replace(hour=23, minute=59, second=59, microsecond=0).timestamp() * 1000))
    return str(int((now.timestamp() + int(sel) * 3600) * 1000))


def publish(deck_dir, expiry='72', dry_run=False):
    """Present-only copy at <site>/d/<id>.<epochMs|x>.html. The cleanup action deletes it after expiry."""
    ep = expiry_epoch(expiry)
    page = inject(single_html(deck_dir), head=f'<meta name="deck-expires" content="{"never" if ep == "x" else ep}">\n'
                  f'<script>{GATE}{CLOSE_SCRIPT}')
    page = re.sub(r'<aside class="notes">.*?</aside>', '', page, flags=re.S)   # speaker notes stay private
    name = f'{os.urandom(6).hex()}.{ep}.html'
    payload = os.path.join(tempfile.mkdtemp(prefix='deck-pub-'), 'payload.json')
    json.dump({'message': f'publish {name}', 'branch': GH['branch'],
               'content': base64.b64encode(page.encode('utf-8')).decode()}, open(payload, 'w'))
    url = f'{GH["site"]}/d/{name}'
    if dry_run:
        return url, payload
    try:
        gh('-X', 'PUT', f'repos/{GH["owner"]}/{GH["repo"]}/contents/d/{name}', input_file=payload)
    finally:
        shutil.rmtree(os.path.dirname(payload), ignore_errors=True)
    return url


def published():
    try:
        items = gh(f'repos/{GH["owner"]}/{GH["repo"]}/contents/d?ref={GH["branch"]}')
    except RuntimeError as e:
        if '404' in str(e):
            return []
        raise
    out = []
    for f in items or []:
        if f['type'] != 'file' or not f['name'].endswith('.html'):
            continue
        ep = f['name'].split('.')[-2]
        exp = None if ep == 'x' else datetime.datetime.fromtimestamp(int(ep) / 1000).isoformat(timespec='minutes')
        out.append({'name': f['name'], 'url': f'{GH["site"]}/d/{f["name"]}', 'expires': exp, 'sha': f['sha']})
    return out


def revoke(name):
    for f in published():
        if f['name'] == name:
            gh('-X', 'DELETE', f'repos/{GH["owner"]}/{GH["repo"]}/contents/d/{name}',
               '-f', f'message=revoke {name}', '-f', f'sha={f["sha"]}', '-f', f'branch={GH["branch"]}')
            return True
    return False


def main():
    a = sys.argv[1:]
    if a[:1] == ['_pptx']:
        return write_pptx(a[1], a[2])
    if a[:1] == ['published']:
        for f in published():
            print(f'{f["name"]:32} {f["expires"] or "never":18} {f["url"]}')
        return
    if a[:1] == ['revoke']:
        print('revoked' if revoke(a[1]) else 'not found')
        return
    if a[:1] == ['publish']:
        r = publish(a[1], a[2] if len(a) > 2 and not a[2].startswith('-') else '72', dry_run='--dry-run' in a)
        print(*(r if isinstance(r, tuple) else (r,)))
        return
    import argparse
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('deck'); ap.add_argument('--formats', default='html,pdf,pptx'); ap.add_argument('--dest')
    x = ap.parse_args(a)
    for p in export(x.deck, [f.strip() for f in x.formats.split(',') if f.strip()], x.dest):
        print(p)


if __name__ == '__main__':
    main()
